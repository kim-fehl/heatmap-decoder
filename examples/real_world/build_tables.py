"""Rebuild source tables with `uv run --with openpyxl python .../build_tables.py`."""

import csv
import hashlib
import json
from pathlib import Path

import numpy as np
import openpyxl
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent


def write_tsv(path, header, rows):
    with path.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.writer(stream, delimiter="\t", lineterminator="\n")
        writer.writerow(header)
        writer.writerows(rows)


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    source = ROOT / "sources/elife-78496/elife-78496-fig5-data1-v2.xlsx"
    sheet = openpyxl.load_workbook(source, data_only=True)["Figure 5E"]
    out = ROOT / "elife-78496-fig5e"
    genes = (out / "rows.txt").read_text().splitlines()
    # Empty spacer columns G, L and P do not represent samples.
    source_columns = [3, 4, 5, 6, 8, 9, 10, 11, 13, 14, 15, 17, 18, 19]
    groups = [
        ("BMAd-Pnpla2+/+", "Ad libitum", 4),
        ("BMAd-Pnpla2+/+", "CR", 4),
        ("BMAd-Pnpla2-/-", "Ad libitum", 3),
        ("BMAd-Pnpla2-/-", "CR", 3),
    ]
    columns, mappings = [], []
    i = 0
    for genotype, diet, count in groups:
        for ordinal in range(1, count + 1):
            label = f"{genotype}__{diet}__{ordinal:02d}"
            columns.append(label)
            mappings.append(
                [
                    i,
                    label,
                    genotype,
                    diet,
                    ordinal,
                    "display_position_within_group",
                    sheet.title,
                    get_column_letter(source_columns[i]),
                ]
            )
            i += 1
    source_rows = {}
    for row in range(3, sheet.max_row + 1):
        gene = sheet.cell(row, 2).value
        if gene in source_rows:
            raise ValueError(f"Duplicate source gene: {gene}")
        source_rows[gene] = row
    assert len(set(genes)) == len(genes) == 49
    assert set(genes) == set(source_rows)
    truth = np.asarray(
        [[sheet.cell(source_rows[gene], col).value for col in source_columns] for gene in genes],
        dtype=float,
    )
    assert truth.shape == (49, 14) and np.isfinite(truth).all()
    # The workbook supplies plotted z-scores directly. Do not rescale or reorder samples.
    (out / "columns.txt").write_text("\n".join(columns) + "\n")
    write_tsv(
        out / "truth.tsv",
        ["gene", *columns],
        [[gene, *values] for gene, values in zip(genes, truth, strict=True)],
    )
    write_tsv(
        out / "row_mapping.tsv",
        ["display_index", "display_label", "source_id", "sheet", "source_row"],
        [[i, gene, gene, sheet.title, source_rows[gene]] for i, gene in enumerate(genes)],
    )
    write_tsv(
        out / "column_mapping.tsv",
        [
            "display_index",
            "display_label",
            "genotype",
            "diet",
            "group_ordinal",
            "ordinal_kind",
            "sheet",
            "source_column",
        ],
        mappings,
    )
    manifest = {
        "doi": "10.7554/eLife.78496",
        "panel": "Figure 5E",
        "attribution": "Li et al., eLife 11:e78496 (2022)",
        "license": "CC BY",
        "source_url": "https://cdn.elifesciences.org/articles/78496/elife-78496-fig5-data1-v2.xlsx",
        "image_origin": "User-supplied crop named elife-78496-fig5-v2.jpg",
        "source_workbook_sha256": digest(source),
        "image_sha256": digest(out / "figure.jpg"),
        "truth_sha256": digest(out / "truth.tsv"),
        "shape": list(truth.shape),
        "value_type": "author_supplied_plotted_z_score",
        "additional_transformation": "none",
        "row_order": "visually transcribed top to bottom; exact symbol join to Figure 5E sheet",
        "column_order": "left to right, C:F H:K M:O Q:S; 4/4/3/3 samples",
        "sample_identity": "original sample IDs absent; names use group plus display ordinal",
        "row_order_verified": True,
        "column_group_order_verified": True,
        "column_within_group_order_check": (
            "color agreement to supplied table; no sample IDs visible"
        ),
        "truth_range": [float(truth.min()), float(truth.max())],
    }
    (out / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n")

    source = ROOT / "sources/elife-17290/elife-17290-fig2-data2-v1.xlsx"
    sheet = openpyxl.load_workbook(source, data_only=True).active
    out = ROOT / "elife-17290-fig2"
    records = []
    for row in range(2, sheet.max_row + 1):
        values = [sheet.cell(row, col).value for col in range(1, 6)]
        if all(value is None for value in values):
            continue
        assert isinstance(values[0], str) and isinstance(values[1], str)
        assert np.isfinite(np.asarray(values[2:], dtype=float)).all()
        records.append([sheet.title, row, *values])
    assert len(records) == 598
    write_tsv(
        out / "source_table.tsv",
        [
            "sheet",
            "source_row",
            "sgRNA_sequence",
            "gene",
            "genome_wide_GFP_SQSTM1_high_vs_low_log2FC",
            "mini_pool_GFP_SQSTM1_high_vs_low_log2FC",
            "mini_pool_SQSTM1_high_vs_low_log2FC",
        ],
        records,
    )
    (out / "manifest.json").write_text(
        json.dumps(
            {
                "doi": "10.7554/eLife.17290",
                "attribution": "DeJesus et al., eLife 5:e17290 (2016)",
                "license": "CC BY",
                "source_workbook_sha256": digest(source),
                "source_url": "https://cdn.elifesciences.org/articles/17290/elife-17290-fig2-data2-v1.xlsx",
                "source_table_sha256": digest(out / "source_table.tsv"),
                "records": len(records),
                "status": "panel_G_image_and_exact_guide_records_attached",
                "image_sha256": digest(out / "figure.jpg"),
                "source_direction": "high versus low",
                "row_order_verified": True,
                "note": (
                    "Retains sgRNA and gene keys, source order and signs; no guide aggregation."
                ),
            },
            indent=2,
        )
        + "\n"
    )
    build_panel_g(records)
    build_blood_source()
    build_workbook_inventory()
    print("Rebuilt publication tables, panel G truth and blood source export.")


def build_panel_g(records):
    out = ROOT / "elife-17290-fig2"
    genes = (out / "rows.txt").read_text().splitlines()
    columns = (out / "columns.txt").read_text().splitlines()
    # Source rows 2:21 match every displayed gene, including the two LC3B names.
    selected = records[:20]
    assert [r[3] for r in selected] == genes
    write_tsv(out / "truth.tsv", ["gene", *columns], [[r[3], *r[4:]] for r in selected])
    write_tsv(
        out / "row_mapping.tsv",
        ["display_index", "gene", "source_row", "sgRNA_sequence"],
        [[i, r[3], r[1], r[2]] for i, r in enumerate(selected)],
    )
    manifest = json.loads((out / "manifest.json").read_text())
    manifest.update(
        panel="Figure 2G",
        shape=[20, 3],
        truth_sha256=digest(out / "truth.tsv"),
        value_type="author_supplied_guide_log2FC",
        additional_transformation="none",
        calibration_note=(
            "0 and 6 assigned to visible bar endpoints; endpoint positions are manually estimated"
        ),
        shared_guide_note=(
            "MAP1LC3B and MAP1LC3B2 share a guide and identical source values; retained separately"
        ),
    )
    (out / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n")


def build_blood_source():
    source = ROOT / "elife-97830-fig4-data1-v1.xlsx"
    workbook = openpyxl.load_workbook(source, data_only=True, read_only=True)
    sheet = workbook["DEGSEhighvsLow"]
    records = list(sheet.values)
    header = records[0]
    samples = list(header[9:])
    assert len(samples) == len(set(samples)) == 39
    out = ROOT / "elife-97830-fig4e"
    genes = (out / "rows.txt").read_text().splitlines()
    selected = []
    mapping = []
    for gene in genes:
        matches = [(i, r) for i, r in enumerate(records[1:], 2) if r[1] == gene]
        assert len(matches) == 1, (gene, len(matches))
        source_row, record = matches[0]
        selected.append([gene, *record[9:]])
        mapping.append([gene, record[0], sheet.title, source_row])
    values = np.array([r[1:] for r in selected], dtype=float)
    assert values.shape == (25, 39) and np.isfinite(values).all()
    write_tsv(out / "source_expression.tsv", ["gene", *samples], selected)
    write_tsv(out / "source_row_mapping.tsv", ["gene", "ENSID", "sheet", "source_row"], mapping)
    groups = {label: group for group, label in list(workbook["labelcode"].values)[1:]}
    assert set(groups) == set(samples)
    write_tsv(
        out / "source_samples.tsv",
        ["sample", "group", "source_column"],
        [[sample, groups[sample], get_column_letter(i)] for i, sample in enumerate(samples, 10)],
    )
    manifest = {
        "doi": "10.7554/eLife.97830.3",
        "panel": "Figure 4E",
        "attribution": "Bæk/Muk et al., eLife 13:RP97830",
        "license": "CC BY",
        "source_workbook": source.name,
        "source_workbook_sha256": digest(source),
        "source_expression_sha256": digest(out / "source_expression.tsv"),
        "shape": [25, 39],
        "display_shape": [25, 26],
        "status": "source_expression_available_display_transform_and_order_unverified",
        "row_order_verified": True,
        "column_order_verified": False,
        "additional_transformation": "none; source expression is not plotted z-score truth",
        "note": (
            "All 25 genes matched uniquely; 13 HIGH_SE and 14 LOW_SE samples available, "
            "but image shows 13 per group. Low-group selection, within-group image order "
            "and plotted transformation are not supplied by the workbook. "
            "No color-fitted ordering is treated as independent truth."
        ),
    }
    (out / "source_manifest.json").write_text(json.dumps(manifest, indent=2) + "\n")


def build_workbook_inventory():
    inventory = []
    for source in sorted(ROOT.glob("*.xlsx")):
        workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
        inventory.append(
            {
                "path": source.name,
                "sha256": digest(source),
                "sheets": [
                    {"name": s.title, "rows": s.max_row, "columns": s.max_column} for s in workbook
                ],
            }
        )
        workbook.close()
    (ROOT / "workbooks.json").write_text(json.dumps(inventory, indent=2) + "\n")


if __name__ == "__main__":
    main()
